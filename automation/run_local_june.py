"""Local verification harness (June-2026) for the Pivot All Region pipeline.

Loads the same inputs the Colab notebook would, but from local paths, and:
  1. builds the Data sheet and reconciles aggregates vs the golden workbook,
  2. builds Pivot + Commission with the real June KPI intake,
  3. writes the three sheets into a copy of the workbook (unless --no-write),
  4. diffs the pasted-KPI columns vs golden (by DPID) to prove the KPI wiring.

The authoritative calc oracle stays verify_commission_engine.py (golden inputs
-> outputs, 0 mismatch). Here the pivot-driven columns come from a FRESH Results
run, so they are expected to differ from the June deliverable.
"""
import os
import sys
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.utils import column_index_from_string as ci

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import par_pipeline as par
import par_writer

BASE = r"C:/Users/mohammad.wirawan_nin/Documents/Mitra All Region"
DATA = os.path.join(BASE, "sample-data-june-2026")   # June test inputs
REF = os.path.join(BASE, "reference")                # golden pivots
RESULTS = os.path.join(BASE, "Results", "Shipper ID Level")
GOLDEN = os.path.join(REF, "M06Y2026 Pivot All Region.xlsx")
REGION_SPLIT = r"C:/Users/mohammad.wirawan_nin/Documents/Mitra/Resources/3. Mitra Region Split.xlsx"

# Local June KPI files (in sample-data-june-2026/).
ACQ_FILE = os.path.join(DATA, "Acquisition Jun 26.csv")
ALLO_FILE = os.path.join(DATA, "Allocation KPI Compile - Jun 2026 (1).csv")
UNREG_FILE = os.path.join(DATA, "Unreg KPI Compile - Jun 2026.csv")

SCRATCH = os.environ.get("PAR_OUT_DIR", os.path.join(BASE, "_generated"))
OUT_FILE = os.path.join(SCRATCH, "M06Y2026 Pivot All Region_GENERATED.xlsx")


def load_shipper():
    p = lambda n: os.path.join(RESULTS, n)
    return {
        "akus_alo_unreg": pd.read_csv(p("Shipper Level Akuisisi, Alokasi, dan Unregistered Parcel.csv")),
        "regular_cod":    pd.read_csv(p("Shipper Level Regular & COD Parcel.csv")),
        "last_mile":      pd.read_csv(p("Shipper Level Last Mile Parcel.csv")),
        "lex":            pd.read_csv(p("Shipper Level Lex Hub Pickup Parcel.csv")),
        "mpm":            pd.read_csv(p("Shipper Level Mitra Pickup Mitra.csv")),
    }


def load_region_split():
    df = pd.read_excel(REGION_SPLIT, sheet_name="Sheet1")
    return df.rename(columns={c: c.strip() for c in df.columns})


def _opt_csv(path, label):
    if os.path.exists(path):
        print(f"  [ok] {label}: {os.path.basename(path)}")
        return pd.read_csv(path)
    print(f"  [--] {label}: absent (tolerant mode)")
    return None


def agg(df, label):
    g = df.groupby("Mapping Type").agg(
        rows=("Mapping Type", "size"),
        parcels=("Total Parcel", "sum"),
        weight=("Total Weight", "sum"),
        delfee=("Total Delivery Fee", "sum"),
        comm=("Total Commission", "sum"),
        do_parcel=("Drop off to SH Parcel", "sum"),
        do_comm=("Drop off to SH Commission", "sum"),
        ins=("Insurance Fee", "sum"),
    ).round(2)
    print(f"\n===== {label} =====")
    print(g.to_string())
    return g


# --------------------------------------------------------------------------
# KPI wiring diff (by DPID) vs golden Commission
# --------------------------------------------------------------------------
KPI_CHECK_COLS = ["U", "V", "W", "X", "Y", "Z", "AA", "AB", "AD", "AE",
                  "AP", "AQ", "AR", "AS", "BD", "BE", "BF", "BG", "BJ", "BQ"]

# Which KPI source each checked column derives from (for absent-vs-differs split).
_COL_SOURCE = {**{c: "acq" for c in ["U", "V", "W", "X", "Y", "Z", "AA", "AB", "AD", "AE"]},
               **{c: "allo" for c in ["AP", "AQ", "AR", "AS"]},
               **{c: "unreg" for c in ["BD", "BE", "BF", "BG"]},
               **{c: "pivot" for c in ["BJ", "BQ"]}}


def load_golden_commission_by_dpid():
    wb = openpyxl.load_workbook(GOLDEN, data_only=True)
    ws = wb["Commission"]
    out = {}
    for r in range(4, 173 + 1):
        a = ws.cell(row=r, column=1).value
        if a in (None, ""):
            continue
        out[par._norm_one(a)] = {L: ws.cell(row=r, column=ci(L)).value for L in KPI_CHECK_COLS}
    return out


def diff_kpi(commission, golden_by_dpid, src_keys):
    """Compare pasted-KPI columns vs golden by DPID, splitting each mismatch into
    'absent' (DPID not in our local KPI source -> coverage gap) vs 'differs'
    (DPID present but value differs -> data-version drift / golden fill-down)."""
    ours = {par._norm_one(row["A"]): row for _, row in commission.iterrows()}
    shared = [k for k in ours if k in golden_by_dpid]
    print(f"\n===== KPI wiring diff vs golden (by DPID) — {len(shared)} shared DPIDs =====")
    print("  'absent' = DPID not in our local KPI file (partial); 'differs' = value drift.")
    print(f"{'col':>3} {'match':>6} {'absent':>7} {'differs':>8}  note")
    for L in KPI_CHECK_COLS:
        src = _COL_SOURCE[L]
        keys = src_keys.get(src)
        match = absent = differs = 0
        samples = []
        for k in shared:
            gv, cv = golden_by_dpid[k].get(L), ours[k].get(L)
            if L == "AB":
                ok = str(gv).strip() == str(cv).strip()
            else:
                gvn, cvn = par._num(gv), par._num(cv)
                ok = abs(gvn - cvn) <= 1e-6 + 1e-4 * max(abs(gvn), abs(cvn))
            if ok:
                match += 1
            elif keys is not None and k not in keys:
                absent += 1
            else:
                differs += 1
                if len(samples) < 2:
                    samples.append((k, gv, cv))
        note = ""
        if L in ("U", "Z") and differs:
            note = "golden manual fill-down block (§6) — automation fixes"
        elif samples:
            note = "e.g. " + "; ".join(f"dp{k}: gold={g} got={c}" for k, g, c in samples)
        print(f"{L:>3} {match:>6} {absent:>7} {differs:>8}  {note}")


def main():
    do_write = "--no-write" not in sys.argv

    shipper = load_shipper()
    df_reg_type = load_region_split()
    data = par.build_data_sheet(shipper, df_reg_type)

    print("Built Data rows:", len(data))
    print("Region counts:", data["Region"].value_counts(dropna=False).to_dict())
    mine = agg(data, "GENERATED Data")

    golden_data = pd.read_excel(GOLDEN, sheet_name="Data")
    ggen = agg(golden_data, "GOLDEN Data (M06Y2026)")

    base_types = ["ACQUISITION", "ALLOCATION", "Unregistered Eligible", "Last Mile",
                  "LEX Hub Pickup", "Mitra Pickup Mitra", "Regular", "COD"]
    print("\n===== DELTA (generated - golden) on shared base types =====")
    for t in base_types:
        if t in mine.index and t in ggen.index:
            dm, dg = mine.loc[t], ggen.loc[t]
            print(f"{t:22} rows {int(dm['rows'])-int(dg['rows']):+6d} | "
                  f"parcels {dm['parcels']-dg['parcels']:+.0f} | comm {dm['comm']-dg['comm']:+.0f}")
        else:
            print(f"{t:22} present mine={t in mine.index} golden={t in ggen.index}")

    # ---- Commission (Phase 1 end-to-end) ----
    print("\n===== Loading KPI intake =====")
    df_acq = _opt_csv(ACQ_FILE, "Acquisition KPI")
    df_allo = _opt_csv(ALLO_FILE, "Allocation KPI")
    df_unreg = _opt_csv(UNREG_FILE, "Unreg KPI")
    # df_lipoh / df_lm_tracker / prev pivot: absent locally -> tolerant mode.

    kpi_fn = par.build_kpi_row_fn(df_acq=df_acq, df_allo=df_allo, df_unreg=df_unreg)
    commission = par.assemble_commission(data, kpi_fn)
    print(f"\nBuilt Commission rows: {len(commission)} (expected ~169)")

    # ---- Phase 2b: Commission sort + Rekap / By-Owner key lists ----
    byo = par._is_by_owner(commission["D"])
    print(f"Commission last 8 DPIDs: {commission['A'].tail(8).tolist()}")
    print(f"'By Owner' rows (should be at the bottom): {commission.loc[byo, 'A'].tolist()}")
    rekap_keys = par.build_rekap_keys(commission)
    byowner_keys = par.build_byowner_keys(commission)
    print(f"Rekap keys: {len(rekap_keys)} (last 6: {rekap_keys[-6:]})")
    print(f"By Owner keys: {byowner_keys}")

    src_keys = {
        "acq": set(par._norm_id(df_acq["DP ID"])) if df_acq is not None else None,
        "allo": set(par._norm_id(df_allo["DP ID"])) if df_allo is not None else None,
        "unreg": set(par._norm_id(df_unreg["DP ID"])) if df_unreg is not None else None,
        "pivot": None,
    }
    diff_kpi(commission, load_golden_commission_by_dpid(), src_keys)

    if do_write:
        print("\n===== Writing workbook (LIVE template round-trip; ~3 min) =====")
        par_writer.write_pivot_all_region(GOLDEN, OUT_FILE, data, commission,
                                          rekap_keys=rekap_keys, byowner_keys=byowner_keys,
                                          month="2026-06")
        print(f"[done] wrote {OUT_FILE}")
    else:
        print("\n(skipped workbook write: --no-write)")


if __name__ == "__main__":
    main()
