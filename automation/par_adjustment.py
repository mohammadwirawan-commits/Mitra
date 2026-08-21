"""Phase 2a — Adjustment & Penalty sheet builder (pandas only, openpyxl-light reader).

Reproduces the golden `Adjustment&Penalty` sheet of `M##Y#### Pivot All Region.xlsx`.

Two intake artifacts feed it (mirrors the manual workflow exactly):
  1. the `Adjustment` intake tab  -> the 72 approved-adjustment rows (E/F/G/J);
  2. `Format SR Penalty - {Mon}.xlsx` (the worksheet the team compiles from several
     external files) -> the penalty rows (H/I) and the Penalty Outstanding block (L:Y).

Layout produced (verified against golden M06Y2026):
  main table  A:J  = 72 adjustment rows (E/F/G/J, H/I=0) then N penalty rows
                     (H/I, E/F/G blank), penalty rows in Format SR `Pivot Parcel` order;
  outstanding L:Y  = Format SR `Penalty Outstanding` A:J -> L:U, plus helper cols V/W/X/Y;
  pivots      AA:AF (by DPMS), AH:AM (by DPID), AO:AR (Outstanding by DPMS).

B (DPMS) and D (Region) come from a VLOOKUP into the Data sheet by DP ID, yielding
the literal string '#N/A' when the DP has no commission row this month (golden behaviour).
C (Nama Mitra) comes from the source tab. Deduction split is by the `Case Pivot` column:
  "Penalty Parcel" -> H (Fake POD + Fake PODA + Missing LM), "Penalty PPDP" -> I.
"""
from __future__ import annotations
import openpyxl
import pandas as pd

import par_pipeline as par

NA = "#N/A"                       # Excel VLOOKUP miss (kept as a literal, like golden)
_ID_MONTHS = ["Januari", "Februari", "Maret", "April", "Mei", "Juni",
              "Juli", "Agustus", "September", "Oktober", "November", "Desember"]

# Main-table columns in golden order (A..J).
MAIN_COLS = ["DP ID", "DPMS ID", "Nama Mitra", "Region", "Adjustment Komisi",
             "Adjustment Bruto", "Adjustment PPh", "Penalty Parcel", "Penalty PPDP", "Note"]
# Outstanding block columns in golden order (L..Y).
OS_COLS = ["DP ID", "DPMS ID", "Mitra", "region", "Invoice", "Total Outstanding",
           "Denda", "Total Penalty", "Periode", "Note",
           "OS Gsheet Inv & Comm", "Invoice Periode", "Bulan", "Periode2"]


# --------------------------------------------------------------------------- readers
def _locate_header(ws, needles, max_scan=12):
    """Return the 1-based row index whose cells contain all `needles` (case-insensitive)."""
    want = {n.strip().lower() for n in needles}
    for r in range(1, min(ws.max_row, max_scan) + 1):
        have = {str(ws.cell(r, c).value).strip().lower()
                for c in range(1, ws.max_column + 1) if ws.cell(r, c).value is not None}
        if want <= have:
            return r
    raise ValueError(f"header row with {needles} not found in {ws.title}")


def _sheet_records(ws, header_row):
    """Rows below `header_row` as dicts keyed by header label (blank-key cols dropped)."""
    hdr = [ws.cell(header_row, c).value for c in range(1, ws.max_column + 1)]
    keep = [(c, h) for c, h in enumerate(hdr, start=1) if h not in (None, "")]
    out = []
    for r in range(header_row + 1, ws.max_row + 1):
        rec = {h: ws.cell(r, c).value for c, h in keep}
        if all(v is None for v in rec.values()):
            continue
        out.append(rec)
    return out


def read_format_sr(path):
    """Read the Format SR Penalty workbook -> (penalty_detail, pivot_order, outstanding).

    penalty_detail : list of dicts from the `Penalty Parcel` sheet (tracking level).
    pivot_order    : list of DP-ID keys in the `Pivot Parcel` sheet order (or None).
    outstanding    : list of dicts from the `Penalty Outstanding` sheet.
    """
    wb = openpyxl.load_workbook(path, data_only=True)

    ws = wb["Penalty Parcel"]
    hr = _locate_header(ws, ["DP ID", "Case Pivot", "Deduction Amount"])
    detail = _sheet_records(ws, hr)

    order = None
    if "Pivot Parcel" in wb.sheetnames:
        pv = wb["Pivot Parcel"]
        try:
            phr = _locate_header(pv, ["DP ID", "Penalty Parcel", "Penalty PPDP"])
            order = [par._norm_one(rec["DP ID"]) for rec in _sheet_records(pv, phr)
                     if rec.get("DP ID") is not None
                     and "grand" not in str(rec.get("DP ID")).lower()]
        except ValueError:
            order = None

    os_rows = []
    if "Penalty Outstanding" in wb.sheetnames:
        osw = wb["Penalty Outstanding"]
        ohr = _locate_header(osw, ["DP ID", "Total Outstanding", "Total Penalty"])
        os_rows = _sheet_records(osw, ohr)

    return detail, order, os_rows


# --------------------------------------------------------------------------- builder
def _data_lookup(data: pd.DataFrame):
    """DP-ID key -> (DPMS, Region) from the built Data sheet (first occurrence wins)."""
    dpms, region = {}, {}
    for _, r in data.iterrows():
        k = par._norm_one(r["DPID"])
        if k and k not in dpms:
            dpms[k] = r["DPMS"]
            region[k] = r["Region"]
    return dpms, region


def _invoice_suffix(inv):
    """Last '-NN' of an invoice string like '2025-18616-12' -> '-12'."""
    s = str(inv) if inv is not None else ""
    parts = s.rsplit("-", 1)
    return "-" + parts[1] if len(parts) == 2 else ""


def build_adjustment_penalty(data: pd.DataFrame, df_adjustment: pd.DataFrame | None,
                             format_sr_path: str | None = None,
                             penalty_detail: pd.DataFrame | None = None,
                             outstanding_path: str | None = None,
                             outstanding_df: pd.DataFrame | None = None,
                             verbose: bool = True) -> dict:
    """Build the Adjustment&Penalty blocks.

    Penalty rows (H/I) come from EITHER `penalty_detail` (a tracking-level frame from
    par_format_sr — the automated compile) OR the `Penalty Parcel` sheet of a compiled
    `format_sr_path` workbook. The outstanding block (L:Y) source, in precedence order:
    `outstanding_df` (the standardized "Penalty Outstanding" intake tab) > `outstanding_path`
    (a compiled Format SR / finance file) > the `Penalty Outstanding` sheet of `format_sr_path`.
    """
    dpms_of, region_of = _data_lookup(data)

    def _dpms(k):   return dpms_of.get(k, NA)
    def _region(k): return region_of.get(k, NA)

    # --- adjustment rows (E/F/G/J) from the intake Adjustment tab -----------------
    adj_rows = []
    if df_adjustment is not None and len(df_adjustment):
        adf = par._clean_cols(df_adjustment.copy())
        c_dp   = par._first_col(adf, ["DP ID", "DPID", "DP Id"])
        c_mit  = par._first_col(adf, ["Nama Mitra", "Mitra Name", "Mitra"])
        c_kom  = par._first_col(adf, ["Adjustment Komisi", "Komisi"])
        c_bru  = par._first_col(adf, ["Adjustment Bruto", "Bruto"])
        c_pph  = par._first_col(adf, ["Adjustment PPh", "PPh"])
        c_note = par._first_col(adf, ["Note", "Notes", "Keterangan"])
        for _, r in adf.iterrows():
            if c_dp is None or pd.isna(r.get(c_dp)):
                continue
            k = par._norm_one(r[c_dp])
            adj_rows.append({
                "DP ID": _numish(r[c_dp]), "DPMS ID": _dpms(k),
                "Nama Mitra": r.get(c_mit), "Region": _region(k),
                "Adjustment Komisi": _numish(r.get(c_kom)),
                "Adjustment Bruto": _numish(r.get(c_bru)),
                "Adjustment PPh": _numish(r.get(c_pph)),
                "Penalty Parcel": 0, "Penalty PPDP": 0,
                "Note": r.get(c_note),
            })

    # --- penalty rows (H/I): compiled detail frame OR compiled Format SR workbook --
    pen_rows, os_records = [], []
    detail_recs, order = None, None
    if penalty_detail is not None:
        detail_recs = penalty_detail.to_dict("records")      # automated compile (par_format_sr)
    elif format_sr_path:
        detail_recs, order, os_records = read_format_sr(format_sr_path)
    # Outstanding source precedence: explicit DataFrame > separate file > compiled Format SR.
    if outstanding_df is not None:                           # standardized "Penalty Outstanding" tab
        os_records = outstanding_df.to_dict("records")
    elif outstanding_path:
        _, _, os_records = read_format_sr(outstanding_path)

    if detail_recs is not None:
        agg = {}                                     # dp key -> dict(H, I, mitra)
        for rec in detail_recs:
            k = par._norm_one(rec.get("DP ID"))
            if not k:
                continue
            cp = str(rec.get("Case Pivot") or "").strip().lower()
            amt = par._num(rec.get("Deduction Amount"))
            slot = agg.setdefault(k, {"H": 0.0, "I": 0.0, "mitra": rec.get("Nama Mitra")})
            if cp == "penalty parcel":
                slot["H"] += amt
            elif cp == "penalty ppdp":
                slot["I"] += amt
        keys = order if order else list(agg.keys())
        seen = set()
        for k in keys:
            if k in seen or k not in agg:
                continue
            seen.add(k)
            v = agg[k]
            if v["H"] == 0 and v["I"] == 0:
                continue
            pen_rows.append({
                "DP ID": _numish(k), "DPMS ID": _dpms(k),
                "Nama Mitra": v["mitra"], "Region": _region(k),
                "Adjustment Komisi": None, "Adjustment Bruto": None, "Adjustment PPh": None,
                "Penalty Parcel": v["H"] or None, "Penalty PPDP": v["I"] or None,
                "Note": "Penalty",
            })

    main = pd.DataFrame(adj_rows + pen_rows, columns=MAIN_COLS)

    # --- Penalty Outstanding block (L:Y) ------------------------------------------
    os_out = []
    for i, rec in enumerate(os_records, start=1):
        get = lambda *names: rec.get(_first_key(rec, names))
        q = par._num(get("Total Outstanding"))
        os_out.append({
            "DP ID": get("DP ID"), "DPMS ID": get("DPMS", "DPMS ID"),
            "Mitra": get("Nama Mitra", "Mitra"), "region": get("region", "Region"),
            "Invoice": get("Invoice"), "Total Outstanding": get("Total Outstanding"),
            "Denda": get("Penalty", "Denda"), "Total Penalty": get("Total Penalty"),
            "Periode": get("Periode"), "Note": get("Note"),
            "OS Gsheet Inv & Comm": -q,
            "Invoice Periode": _invoice_suffix(get("Invoice")),
            "Bulan": f"-{i:02d}",
            "Periode2": _ID_MONTHS[(i - 1) % 12],
        })
    outstanding = pd.DataFrame(os_out, columns=OS_COLS)

    # --- pivots (static value tables) ---------------------------------------------
    pivot_dpms = _pivot_main(main, "DPMS ID")
    pivot_dpid = _pivot_main(main, "DP ID")
    pivot_os = _pivot_os(outstanding)

    if verbose:
        print(f"[adj] adjustment rows={len(adj_rows)}  penalty rows={len(pen_rows)}  "
              f"outstanding rows={len(os_out)}")
        print(f"[adj] totals  E={_col_sum(main,'Adjustment Komisi'):,.1f}  "
              f"F={_col_sum(main,'Adjustment Bruto'):,.1f}  "
              f"H={_col_sum(main,'Penalty Parcel'):,.1f}  "
              f"I={_col_sum(main,'Penalty PPDP'):,.1f}  "
              f"Outstanding={_col_sum(outstanding,'Total Outstanding'):,.1f}")

    return {"main": main, "outstanding": outstanding,
            "pivot_dpms": pivot_dpms, "pivot_dpid": pivot_dpid, "pivot_os": pivot_os}


# --------------------------------------------------------------------------- helpers
def _numish(v):
    """Return an int/float when the value is numeric-looking, else the value as-is."""
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    try:
        f = float(v)
        return int(f) if f.is_integer() else f
    except (TypeError, ValueError):
        return v


def _first_key(rec, names):
    low = {str(k).strip().lower(): k for k in rec}
    for n in names:
        if n.strip().lower() in low:
            return low[n.strip().lower()]
    return None


def _col_sum(df, col):
    return pd.to_numeric(df[col], errors="coerce").fillna(0).sum() if col in df else 0.0


def _pivot_main(main: pd.DataFrame, key: str) -> pd.DataFrame:
    """Group the main table by DPMS/DPID; sum the 5 measures; append a Grand Total row."""
    measures = ["Adjustment Komisi", "Adjustment Bruto", "Adjustment PPh",
                "Penalty Parcel", "Penalty PPDP"]
    d = main.copy()
    d["_label"] = d[key].map(lambda v: par._norm_one(v) if v not in (None, NA) else NA)
    for m in measures:
        d[m] = pd.to_numeric(d[m], errors="coerce").fillna(0.0)
    g = d.groupby("_label", sort=False)[measures].sum().reset_index()
    g["_sort"] = pd.to_numeric(g["_label"], errors="coerce")
    g = g.sort_values(["_sort", "_label"], na_position="last").drop(columns="_sort")
    total = {"_label": "Grand Total", **{m: g[m].sum() for m in measures}}
    out = pd.concat([g, pd.DataFrame([total])], ignore_index=True)
    return out.rename(columns={"_label": "Row Labels"})


def _pivot_os(outstanding: pd.DataFrame) -> pd.DataFrame:
    measures = ["Total Outstanding", "Denda", "Total Penalty"]
    if not len(outstanding):
        return pd.DataFrame(columns=["Row Labels", *measures])
    d = outstanding.copy()
    d["_label"] = d["DPMS ID"].map(lambda v: par._norm_one(v) if v not in (None, NA) else NA)
    for m in measures:
        d[m] = pd.to_numeric(d[m], errors="coerce").fillna(0.0)
    g = d.groupby("_label", sort=False)[measures].sum().reset_index()
    g["_sort"] = pd.to_numeric(g["_label"], errors="coerce")
    g = g.sort_values(["_sort", "_label"], na_position="last").drop(columns="_sort")
    total = {"_label": "Grand Total", **{m: g[m].sum() for m in measures}}
    out = pd.concat([g, pd.DataFrame([total])], ignore_index=True)
    return out.rename(columns={"_label": "Row Labels"})
