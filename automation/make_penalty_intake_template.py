"""Build the standardized "Penalty Intake" template (one Google-Sheet-ready workbook).

Separate from the master intake (PART C decision): the 4 penalty inputs come from different
external stakeholders, so they live in their OWN shared sheet. Each tab is NEAR-RAW — it
mirrors the relevant columns of the stakeholder's real export, so they reformat minimally;
the pipeline (`par_format_sr._compile_detail`) still applies every filter and the hub->DP
mapping. Upload this to Drive, "Open as Google Sheets", share the ONE link; each owner pastes
into their tab; then File -> Download -> .xlsx as "Penalty Intake {Mon-YYYY}.xlsx" into the
month's Intake folder.

Tab names are read by the pipeline (par_format_sr.INTAKE_TAB_*) — do NOT rename them.
Column schemas are the verified Phase 2a+ ones (reconcile to golden M06 exactly).
"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                    "..", "TEMPLATES", "Penalty Intake TEMPLATE.xlsx"))

HDR_FILL = PatternFill("solid", fgColor="305496")
HDR_FONT = Font(bold=True, color="FFFFFF")
EX_FILL = PatternFill("solid", fgColor="F2F2F2")
EX_FONT = Font(italic=True, color="808080")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# tab -> (owner, colour, filter/what-it-drives, [ (column, example) ])
TABS = [
    ("Penalty PPDP", "Desy", "C00000",
     "Pipeline keeps rows where Total Penalty Amount != 0 -> Adjustment&Penalty col I (Penalty PPDP).",
     [("DP ID", 25096), ("Nama Mitra", "Mitra KBM @ Kuwayuhan"), ("Tracking ID", "SHP3149815926"),
      ("Shipper Name", "Shopee (GRIYA SPERPART KEBUMEN)"), ("Granular Status", "Arrived at Sorting Hub"),
      ("Total Penalty Amount", 30000)]),
    ("Penalty Fake POD-PODA", "Arbeta / Opi", "C0338C",
     "Stack POD and PODA rows with a Type flag. Pipeline keeps Penalty > 0 -> col H (Penalty Parcel); "
     "also the source of the Hub Name -> DP ID map used by the Missing LM tab.",
     [("Type (POD/PODA)", "PODA"), ("Hub Name", "DP-SMD-JTMY"), ("DP ID", 19727),
      ("Nama Mitra", "Mitra SMD @ Jatimulya"), ("Tracking ID", "NLIDAP8014406831"),
      ("Granular Status", "Returned to Sender"), ("Penalty", 525)]),
    ("Penalty Missing LM", "Arbeta / Opi", "C0338C",
     "Pipeline keeps Est Claim Amount > 0 AND Final Check CL != 'Take out' -> col H (Penalty Parcel); "
     "DP is derived from Investigating Hub Name (add inactive hubs to HUB_DP_OVERRIDES if flagged).",
     [("Tracking ID", "NLIDAP8015007477"), ("Investigating Hub Name", "DP-SMD-JTMY"),
      ("Order Granular Status", "Cancelled"), ("Type", "SLA BREACH"),
      ("Est Claim Amount", 138000), ("Final Check CL", "Process to claim")]),
    ("Penalty Outstanding", "Finance", "BF9000",
     "Goes straight to the Penalty Outstanding block (L:Y). Finance supplies Denda + Total Penalty "
     "(the 5% denda has per-row exceptions), so the pipeline uses them as-is.",
     [("DP ID", 19060), ("DPMS ID", 19060), ("Nama Mitra", "Mitra PWT @ Kebokura"),
      ("Region", "All Region - Non POH"), ("Invoice", "2026-19060-02"),
      ("Total Outstanding", -571635), ("Denda", -28581.75), ("Total Penalty", -600216.75),
      ("Periode", "Februari 2026"), ("Note", "Potongan outstanding ditambah denda 5% dari CS periode Februari 2026")]),
]


def _add_tab(wb, tab, colour, header, example):
    ws = wb.create_sheet(title=tab)
    ws.sheet_properties.tabColor = colour
    for c, h in enumerate(header, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = max(12, min(34, len(str(h)) + 4))
    for c, v in enumerate(example, start=1):
        cell = ws.cell(row=2, column=c, value=v)
        cell.font = EX_FONT
        cell.fill = EX_FILL
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30


def _build_readme(wb):
    ws = wb.create_sheet(title="README", index=0)
    ws.sheet_properties.tabColor = "000000"
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 82
    ws["A1"] = "MITRA COMMISSION — PENALTY INTAKE (monthly)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Month: ____________          Deadline: by the 3rd working day"
    ws["A2"].font = Font(bold=True, italic=True)
    lines = [
        "",
        "HOW TO USE",
        "1. Find your tab below (tabs are colour-coded by owner).",
        "2. Paste your export UNDER the header row. Keep the column headers exactly as-is.",
        "3. You can paste your FULL export — the pipeline applies the filters shown below.",
        "4. DELETE the grey example row once you've pasted. Do NOT rename tabs or columns.",
        "5. When everyone is done: File → Download → Microsoft Excel (.xlsx),",
        "   name it  \"Penalty Intake {Mon-YYYY}.xlsx\"  and upload it to the month's Intake folder.",
        "",
        "TAB ASSIGNMENTS",
    ]
    r = 3
    for text in lines:
        ws.cell(row=r, column=1, value=text).font = Font(bold=(text in ("HOW TO USE", "TAB ASSIGNMENTS")))
        r += 1
    for c, lbl in enumerate(["Owner", "Tab", "What it drives / how the pipeline uses it"], start=1):
        cell = ws.cell(row=r, column=c, value=lbl)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HDR_FILL
        cell.alignment = Alignment(horizontal="left", vertical="center")
    r += 1
    for tab, owner, colour, drives, _cols in TABS:
        ws.cell(row=r, column=1, value=owner).font = Font(bold=True, color=colour)
        ws.cell(row=r, column=2, value=tab)
        ws.cell(row=r, column=3, value=drives).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 42
        r += 1
    ws.freeze_panes = "A2"


def main():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for tab, _owner, colour, _drives, cols in TABS:
        header = [c for c, _ex in cols]
        example = [ex for _c, ex in cols]
        _add_tab(wb, tab, colour, header, example)
    _build_readme(wb)
    wb.save(OUT)
    print(f"Saved penalty intake template: {OUT}")
    print(f"Tabs ({len(wb.sheetnames)}): {wb.sheetnames}")


if __name__ == "__main__":
    main()
