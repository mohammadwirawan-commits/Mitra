"""Generate the 'Mitra Intake All Region' template workbook.

One tab per external input the Pivot-All-Region stage needs that is NOT already
produced by the first-stage notebook (Results/) or its intake. Ships with a
README, exact headers, and one example row per tab.

Run:  python make_intake_template.py
Out:  Mitra Intake All Region TEMPLATE.xlsx  (in this folder)
"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "Mitra Intake All Region TEMPLATE.xlsx")

HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(bold=True, color="FFFFFF")
EX_FONT = Font(italic=True, color="808080")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# tab -> (phase, feeds, [(header, example)])
TABS = {
    # Pipeline reads 'Total Score' DIRECTLY (authoritative); components kept for
    # transparency. Mirrors the real 'Acquisition Jun 26' export.
    "Acquisition KPI": ("1", "Commission U:Y + Total Score -> payout", [
        ("DP ID", 6857), ("Mitra Name", "Mitra BDO @ Cisondari"),
        ("LI_N0", 0.9318), ("LI_N1", 0.9616), ("PU Scan_N0", 0.9988),
        ("Update_RAR", 1.0), ("PoPA", 0.8801), ("Total Score", 0.9567),
        ("% Payout Commission (Ori)", 1.0), ("Status (Ori)", "Commission As Is"),
        ("Tag", "Non Drop Off")]),
    "Allocation KPI": ("1", "Commission AP:AQ", [
        ("DP ID", 45605), ("Mitra Name", "Mitra BDO @ Mekarrahayu 2"), ("Type", "POH"),
        ("N0_Attempt", 0.998594), ("RoT_Prior_B2B", 1.0)]),
    "Unreg KPI": ("1", "Commission BD:BE", [
        ("DP ID", 45605), ("Mitra Name", "Mitra BDO @ Mekarrahayu 2"), ("Type", "POH"),
        ("LI_N0", 0.827068), ("PU Scan_N0", 0.99306), ("LI_POH", 0.999017)]),
    "Cross Border": ("1", "Data sheet (Cross Border rows; comm = 30% x fee)", [
        ("tracking_id", "SPXXB000000001"), ("DPID", 19727), ("DPMS ID", 19727),
        ("Region Split", "All Region - POH"), ("delivery_fee", 45720)]),
    "Penalty PPDP": ("2", "Adjustment&Penalty col I (Penalty PPDP)", [
        ("DPMS ID", 18616), ("DP Name", "Mitra JKT @ Petojo Selatan"),
        ("Total Penalty Amount", 17450412)]),
    "Penalty Fake POD-PODA": ("2", "Adjustment&Penalty col H (Penalty Parcel)", [
        ("Hub Name", "HUB-JKT-01"), ("Driver Name", "Budi"), ("DPMS ID", 18616),
        ("Type", "PODA"), ("Amount", 250000)]),
    "Penalty Missing LM": ("2", "Adjustment&Penalty col H (Penalty Parcel)", [
        ("Investigasi Hub", "HUB-BDO-02"), ("DPMS ID", 20590),
        ("Estimasi Klaim", 180000)]),
    "Adjustment": ("2", "Adjustment&Penalty cols E:J", [
        ("DP ID", 5707), ("DPMS ID", 5707), ("Nama Mitra", "Mitra JKT @ Balimester 2"),
        ("Region", "All Region - Non POH"), ("Adjustment Komisi", -22208),
        ("Adjustment Bruto", 0), ("Adjustment PPh", 0), ("Note", "Adjustment shopee")]),
    "Penalty Outstanding": ("2", "Adjustment&Penalty cols L:Y + Rekap", [
        ("DP ID", 18616), ("DPMS ID", 18616), ("Mitra", "Mitra JKT @ Petojo Selatan"),
        ("Region", "All Region - Non POH"), ("Invoice", "2025-18616-12"),
        ("Total Outstanding", -23662477), ("Periode", "Desember 2025"),
        ("Note", "Denda 5% dari outstanding")]),
    "Data Mitra Changes": ("3", "Data Sources / Rekap / PPh", [
        ("DPMS ID", 2265), ("DP Name", "Mitra JKT @ Pejuang"), ("Owner Name", "Andi"),
        ("Customer ID Netsuite", 12345), ("Tax ID", "01.234.567.8-901.000"),
        ("Jenis Usaha", "Pribadi"), ("Email", "mitra@example.com"),
        ("Address 1", "Jl. Contoh No.1"), ("Address 2", "Bekasi"),
        ("Bank", "BCA"), ("Account Holder", "Andi"), ("Account Number", "1234567890"),
        ("SKB Number", ""), ("SKB Valid Until", ""), ("Hold Commission", "No"),
        ("Note Perubahan", "")]),
}

README = [
    ("Mitra Intake — All Region", None),
    ("", None),
    ("One tab per external input the 'Pivot All Region' stage needs that is NOT already", None),
    ("produced by the first-stage notebook (Results/) or its intake.", None),
    ("Fill one row per record under each header. Do not rename tabs or headers.", None),
    ("Header row = row 1 (pandas header=0). Grey italic row 2 = example, delete before use.", None),
    ("", None),
    ("Tab | Phase | Feeds", "hdr"),
]


def style_header(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
        ws.column_dimensions[cell.column_letter].width = max(14, len(str(cell.value)) + 3)
    ws.freeze_panes = "A2"


def main():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "README"
    r = 1
    for text, kind in README:
        cell = ws.cell(row=r, column=1, value=text)
        if r == 1:
            cell.font = Font(bold=True, size=14)
        if kind == "hdr":
            cell.font = Font(bold=True)
        r += 1
    for tab, (phase, feeds, _) in TABS.items():
        ws.cell(row=r, column=1, value=f"{tab}  |  Phase {phase}  |  {feeds}")
        r += 1
    ws.column_dimensions["A"].width = 100

    for tab, (phase, feeds, cols) in TABS.items():
        w = wb.create_sheet(tab)
        for i, (h, ex) in enumerate(cols, 1):
            w.cell(row=1, column=i, value=h)
            c = w.cell(row=2, column=i, value=ex)
            c.font = EX_FONT
        style_header(w, len(cols))

    wb.save(OUT)
    print(f"[written] {OUT}")
    print(f"tabs: {', '.join(['README'] + list(TABS))}")


if __name__ == "__main__":
    main()
