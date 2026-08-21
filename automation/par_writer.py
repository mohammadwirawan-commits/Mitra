"""Pivot All Region output writer — LIVE workbook mode.

Template round-trip: load a copy of the prior-month/golden workbook and stamp only
the *source data* as values — the Data rows, the Commission key + pasted-KPI columns,
the Adjustment source tables, the Data Sources master, and the Rekap / By-Owner key
lists. Everything else stays as the workbook built it:

  * the PivotTables (Pivot Data, the Commission!CX rollup, the 3 Adjustment pivots)
    are PRESERVED and set to refresh-on-open, with their cache source ranges rewritten
    to the new data extents;
  * the Commission / Rekap / By-Owner formulas (GETPIVOTDATA, VLOOKUP, computed
    columns) are left intact and filled-down / cleared to match the row counts.

So the delivered workbook recomputes itself the first time it is opened in Excel —
matching the manual process — instead of being flattened to dead numbers.

openpyxl loads the ~4 MB workbook slowly (~85 s); a full round-trip is ~3 min.
It round-trips all 5 pivots / 4 caches / 16 sheets cleanly (verified).
"""
from __future__ import annotations
import os
import re
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string as ci
from openpyxl.formula.translate import Translator

import par_pipeline as par
from par_pipeline import DATA_COLUMNS

DATA_HEADER_ROW = 1          # Data: header row 1, values from row 2
COMMISSION_LABEL_ROW = 3     # Commission: labels row 3, values from row 4
COMMISSION_DATA_START = 4
LAST_COMMISSION_COL = "CV"
COMMISSION_FORMULA_PROBE = ci("E")   # a GETPIVOTDATA column, used to find the formula extent

# --- Adjustment & Penalty sheet layout (verified vs golden M06Y2026) -----------
ADJ_SHEET = "Adjustment&Penalty"
ADJ_DATA_START = 4               # labels row 3, values from row 4 (summary SUBTOTAL row 2)
ADJ_MAIN_COLS = list(range(1, 11))            # A..J
ADJ_OS_COLS = list(range(12, 26))             # L..Y

# --- Phase 2b sheets: keyed formula templates ----------------------------------
DATA_SOURCES_SHEET = "Data Sources"
DATA_SOURCES_HEADER_ROW = 1      # header row 1, numbered row 2, data from row 3
DATA_SOURCES_DATA_START = 3
REKAP_SHEET = "Rekap Commission"
REKAP_DATA_START = 7             # title 1-3, header 4-5, numbered row 6, data row 7
REKAP_PROBE_COL = 2              # col B = VLOOKUP formula
BYOWNER_SHEET = "By Owner per Mitra"
BYOWNER_DATA_START = 8           # header row 5, numbered row 7, data row 8
BYOWNER_PROBE_COL = 2


# --------------------------------------------------------------------------
# small helpers
# --------------------------------------------------------------------------
def _cell_val(v):
    """Coerce a value to something openpyxl can store (no numpy scalars / NaN)."""
    if v is None:
        return None
    if isinstance(v, float) and np.isnan(v):
        return None
    if isinstance(v, np.integer):
        return int(v)
    if isinstance(v, np.floating):
        f = float(v)
        return None if np.isnan(f) else f
    if isinstance(v, float) and v.is_integer():
        return v
    return v


def _set(ws, row, col, value):
    """Assign a cell value. NOTE: ws.cell(row, col, value=None) is a no-op in
    openpyxl (it only assigns when value is not None), so always assign via the
    .value attribute — this is what makes clears (value=None) actually stick."""
    ws.cell(row=row, column=col).value = value


def _is_formula(v) -> bool:
    return isinstance(v, str) and v.startswith("=")


def _find_last_row(ws, r0, key_col=1, max_blank=5):
    last, r, blanks = r0 - 1, r0, 0
    while r <= ws.max_row and blanks < max_blank:
        if ws.cell(row=r, column=key_col).value not in (None, ""):
            last, blanks = r, 0
        else:
            blanks += 1
        r += 1
    return last


def _formula_extent(ws, r0, probe_col):
    """Last row (>= r0-1) for which the probe column still holds a formula — i.e.
    how far the template's formula rows are pre-filled."""
    last, r, blanks = r0 - 1, r0, 0
    while r <= ws.max_row and blanks < 5:
        if _is_formula(ws.cell(row=r, column=probe_col).value):
            last, blanks = r, 0
        else:
            blanks += 1
        r += 1
    return last


def _filldown_formulas(ws, src_row, dst_row, formula_cols):
    """Copy the formula cells of src_row into dst_row, translating relative refs."""
    for c in formula_cols:
        src = ws.cell(row=src_row, column=c)
        if not _is_formula(src.value):
            continue
        origin = f"{get_column_letter(c)}{src_row}"
        dest = f"{get_column_letter(c)}{dst_row}"
        ws.cell(row=dst_row, column=c).value = Translator(src.value, origin=origin).translate_formula(dest)


# --------------------------------------------------------------------------
# top-level writer
# --------------------------------------------------------------------------
def write_pivot_all_region(template_path, out_path, data_df, commission_df,
                           adj_result=None, data_sources_changes=None,
                           rekap_keys=None, byowner_keys=None, verbose=True) -> str:
    """Write the monthly workbook as a LIVE template round-trip (values + preserved
    pivots/formulas). See module docstring.

    `data_df`               16-col Data frame (par.DATA_COLUMNS).
    `commission_df`         one row per DPID, columns keyed by Excel letter (A..CV).
    `adj_result`            optional dict from par_adjustment.build_adjustment_penalty.
    `data_sources_changes`  optional 'Data Mitra Changes' intake frame (upserted).
    `rekap_keys`            optional list of Rekap Commission col-A keys.
    `byowner_keys`          optional list of By Owner per Mitra col-A keys.
    """
    if verbose:
        print(f"[writer] loading template: {template_path}")
    wb = openpyxl.load_workbook(template_path)

    counts = {}
    counts["data"] = _write_data(wb, data_df, verbose)
    counts["commission"] = _write_commission(wb, commission_df, verbose)
    if adj_result is not None:
        counts["adj_main"], counts["adj_os"] = _write_adjustment(wb, adj_result, verbose)
    if data_sources_changes is not None:
        _write_data_sources(wb, data_sources_changes, verbose)
    if rekap_keys is not None:
        _write_keyed_sheet(wb, REKAP_SHEET, REKAP_DATA_START, REKAP_PROBE_COL,
                           rekap_keys, verbose)
    if byowner_keys is not None:
        _write_keyed_sheet(wb, BYOWNER_SHEET, BYOWNER_DATA_START, BYOWNER_PROBE_COL,
                           byowner_keys, verbose)

    _refresh_pivots(wb, counts, verbose)

    if verbose:
        print(f"[writer] saving: {out_path}")
    out_dir = os.path.dirname(out_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    wb.save(out_path)
    return out_path


# --------------------------------------------------------------------------
# Data
# --------------------------------------------------------------------------
def _write_data(wb, data_df, verbose) -> int:
    ws = wb["Data"]
    r0 = DATA_HEADER_ROW + 1
    n = len(data_df)
    ncol = len(DATA_COLUMNS)
    clear_to = max(ws.max_row, r0 + n - 1)
    for r in range(r0, clear_to + 1):               # clear old body first
        for j in range(1, ncol + 1):
            _set(ws, r, j, None)
    for i, (_, row) in enumerate(data_df.iterrows()):
        r = r0 + i
        for j, col in enumerate(DATA_COLUMNS, start=1):
            _set(ws, r, j, _cell_val(row[col]))
    if verbose:
        print(f"[writer] Data: cleared rows {r0}..{clear_to}, wrote {n} rows")
    return n


# --------------------------------------------------------------------------
# Commission — stamp ONLY the value columns; keep every formula column live
# --------------------------------------------------------------------------
def _write_commission(wb, commission_df, verbose) -> int:
    ws = wb["Commission"]
    letters = [get_column_letter(c) for c in range(1, ci(LAST_COMMISSION_COL) + 1)]
    r0 = COMMISSION_DATA_START
    n = len(commission_df)

    # Classify columns by the template's first data row: formula vs value.
    formula_letters = [L for L in letters if _is_formula(ws.cell(row=r0, column=ci(L)).value)]
    formula_cols = [ci(L) for L in formula_letters]
    value_cols = [L for L in letters if L not in formula_letters and L in commission_df.columns]

    tmpl_last = _formula_extent(ws, r0, COMMISSION_FORMULA_PROBE)

    # 1. make sure a formula row exists for every data row (fill down if we grew)
    for i in range(n):
        r = r0 + i
        if r > tmpl_last:
            _filldown_formulas(ws, tmpl_last, r, formula_cols)

    # 2. stamp the value columns for each data row
    for i, (_, row) in enumerate(commission_df.iterrows()):
        r = r0 + i
        for L in value_cols:
            _set(ws, r, ci(L), _cell_val(row.get(L)))

    # 3. clear surplus template rows below the data (formula + value cells)
    last_written = r0 + n - 1
    clear_to = max(tmpl_last, last_written)
    for r in range(last_written + 1, clear_to + 1):
        for L in letters:
            _set(ws, r, ci(L), None)

    if verbose:
        missing = [L for L in letters if L not in formula_letters and L not in commission_df.columns]
        print(f"[writer] Commission: {n} rows (rows {r0}..{last_written}); "
              f"{len(value_cols)} value cols stamped, {len(formula_letters)} formula cols kept; "
              f"template formula extent row {tmpl_last}"
              + (f"; WARNING value cols not produced: {missing}" if missing else ""))
    return n


# --------------------------------------------------------------------------
# Adjustment & Penalty — stamp the source tables; keep the 3 pivots + SUBTOTAL row
# --------------------------------------------------------------------------
def _write_table(ws, df, r0, cols):
    for i, (_, row) in enumerate(df.iterrows()):
        r = r0 + i
        for c, val in zip(cols, row.tolist()):
            _set(ws, r, c, _cell_val(val))
    return r0 + len(df) - 1


def _clear_block(ws, r0, cols, extra_blank_rows=0):
    last = _find_last_row(ws, r0, key_col=cols[0])
    clear_to = last + extra_blank_rows
    for r in range(r0, clear_to + 1):
        for c in cols:
            _set(ws, r, c, None)
    return clear_to


def write_adjustment_penalty(template_path, out_path, adj_result, verbose=True) -> str:
    """Stand-alone round-trip writer for just the Adjustment&Penalty sheet."""
    if verbose:
        print(f"[writer] loading template: {template_path}")
    wb = openpyxl.load_workbook(template_path)
    _write_adjustment(wb, adj_result, verbose)
    _refresh_pivots(wb, {}, verbose)
    if verbose:
        print(f"[writer] saving: {out_path}")
    wb.save(out_path)
    return out_path


def _write_adjustment(wb, res, verbose):
    if ADJ_SHEET not in wb.sheetnames:
        raise KeyError(f"sheet '{ADJ_SHEET}' not in template")
    ws = wb[ADJ_SHEET]

    main = res["main"]
    ost = res["outstanding"]

    # main table A:J and outstanding L:Y are the PIVOT SOURCE tables -> stamp as values.
    # The 3 embedded pivots (AA:AF, AH:AM, AO:AR) and the row-2 SUBTOTAL summary are
    # left intact; they recompute on open (refresh-on-open + fullCalcOnLoad).
    _clear_block(ws, ADJ_DATA_START, ADJ_MAIN_COLS)
    _write_table(ws, main, ADJ_DATA_START, ADJ_MAIN_COLS)

    _clear_block(ws, ADJ_DATA_START, ADJ_OS_COLS)
    _write_table(ws, ost, ADJ_DATA_START, ADJ_OS_COLS)

    if verbose:
        print(f"[writer] Adjustment&Penalty: main={len(main)} rows, outstanding={len(ost)} rows "
              f"(pivots + SUBTOTAL summary left live)")
    return len(main), len(ost)


# --------------------------------------------------------------------------
# Data Sources — upsert the mitra master (values)
# --------------------------------------------------------------------------
def _read_sheet_table(ws, header_row, data_start):
    """Read a sheet's header + body (from data_start) into a DataFrame."""
    headers = []
    for c in range(1, ws.max_column + 1):
        h = ws.cell(row=header_row, column=c).value
        headers.append(h if h not in (None, "") else f"_col{c}")
    rows = []
    last = _find_last_row(ws, data_start, key_col=1)
    for r in range(data_start, last + 1):
        rows.append([ws.cell(row=r, column=c).value for c in range(1, len(headers) + 1)])
    return pd.DataFrame(rows, columns=headers), headers, last


def _write_data_sources(wb, changes, verbose):
    if DATA_SOURCES_SHEET not in wb.sheetnames:
        raise KeyError(f"sheet '{DATA_SOURCES_SHEET}' not in template")
    ws = wb[DATA_SOURCES_SHEET]
    master, headers, old_last = _read_sheet_table(ws, DATA_SOURCES_HEADER_ROW, DATA_SOURCES_DATA_START)

    merged = par.upsert_data_sources(master, changes)

    r0 = DATA_SOURCES_DATA_START
    n = len(merged)
    clear_to = max(old_last, r0 + n - 1)
    for r in range(r0, clear_to + 1):
        for c in range(1, len(headers) + 1):
            _set(ws, r, c, None)
    for i, (_, row) in enumerate(merged.iterrows()):
        r = r0 + i
        for c, col in enumerate(headers, start=1):
            _set(ws, r, c, _cell_val(row.get(col)))
    if verbose:
        added = n - len(master)
        print(f"[writer] Data Sources: {n} mitra rows ({added:+d} vs carried-forward {len(master)})")
    return n


# --------------------------------------------------------------------------
# Rekap Commission / By Owner per Mitra — stamp the col-A key list; keep formulas
# --------------------------------------------------------------------------
def _write_keyed_sheet(wb, sheet, data_start, probe_col, keys, verbose):
    if sheet not in wb.sheetnames:
        raise KeyError(f"sheet '{sheet}' not in template")
    ws = wb[sheet]
    r0 = data_start
    n = len(keys)
    ncol = ws.max_column
    tmpl_last = _formula_extent(ws, r0, probe_col)
    formula_cols = [c for c in range(1, ncol + 1)
                    if _is_formula(ws.cell(row=r0, column=c).value)]

    # fill down formula rows if the key list is longer than the template extent
    for i in range(n):
        r = r0 + i
        if r > tmpl_last:
            _filldown_formulas(ws, tmpl_last, r, formula_cols)

    # stamp the key column (A) for each key
    for i, k in enumerate(keys):
        _set(ws, r0 + i, 1, _cell_val(k))

    # clear surplus template rows below the data
    last_written = r0 + n - 1
    clear_to = max(tmpl_last, last_written)
    for r in range(last_written + 1, clear_to + 1):
        for c in range(1, ncol + 1):
            _set(ws, r, c, None)

    if verbose:
        print(f"[writer] {sheet}: {n} keys (rows {r0}..{last_written}); "
              f"{len(formula_cols)} formula cols kept; template extent row {tmpl_last}")
    return n


# --------------------------------------------------------------------------
# Pivots — refresh on open + rewrite cache source ranges to the new extents
# --------------------------------------------------------------------------
def _reref(ref, nrows):
    """Given a worksheet-source ref 'A1:P16345' and the new DATA row count, keep the
    start cell + end column but move the end row to (start_row + nrows)."""
    m = re.match(r"^([A-Za-z]+)(\d+):([A-Za-z]+)(\d+)$", ref or "")
    if not m:
        return ref
    sc, sr, ec, _er = m.groups()
    return f"{sc}{sr}:{ec}{int(sr) + int(nrows)}"


def _refresh_pivots(wb, counts, verbose=True):
    """Set every pivot cache to refresh-on-open and rewrite its source range to the
    new data extent; force a full recalc on open so the GETPIVOTDATA/VLOOKUP/computed
    formulas recompute against the fresh pivots."""
    touched = 0
    for ws in wb.worksheets:
        for pt in getattr(ws, "_pivots", []) or []:
            cache = pt.cache
            cache.refreshOnLoad = True
            src = getattr(cache.cacheSource, "worksheetSource", None)
            if src is None or not src.ref:
                touched += 1
                continue
            sheet = (src.sheet or "").strip()
            nrows = None
            if sheet == "Data":
                nrows = counts.get("data")
            elif sheet == "Commission":
                nrows = counts.get("commission")
            elif sheet == ADJ_SHEET:
                nrows = counts.get("adj_main") if src.ref[:1].upper() == "A" else counts.get("adj_os")
            if nrows is not None:
                src.ref = _reref(src.ref, nrows)
            touched += 1
    try:
        wb.calculation.fullCalcOnLoad = True
    except Exception:
        pass
    if verbose:
        print(f"[writer] pivots: {touched} set refresh-on-open; fullCalcOnLoad=True")
