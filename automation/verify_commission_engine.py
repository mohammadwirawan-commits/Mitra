"""Verify compute_commission_derived() against the golden June workbook.

Feeds the engine the golden file's own INPUT columns (pivot values + pasted
KPI) and compares the engine's computed FORMULA columns to the golden values.
This isolates the calculation logic from any data-run drift.
"""
import os, sys
import numpy as np
import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string as ci

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import par_pipeline as par

GOLDEN = r"C:/Users/mohammad.wirawan_nin/Documents/Mitra All Region/reference/M06Y2026 Pivot All Region.xlsx"

COMPUTED = ["P","Z","AA","AB","AD","AE","AK","AR","AS","AY","BF","BG","BJ","BQ","BW","CE","CF","CG","CH"]
DATA_ROWS = range(4, 173)  # 169 DPID rows


def main():
    wb = openpyxl.load_workbook(GOLDEN, data_only=True)
    ws = wb["Commission"]
    all_letters = [get_column_letter(c) for c in range(1, ci("CV") + 1)]

    n = 0
    fails = {c: [] for c in COMPUTED}
    for r in DATA_ROWS:
        row = {L: ws.cell(row=r, column=ci(L)).value for L in all_letters}
        if row.get("A") in (None, ""):
            continue
        n += 1
        got = par.compute_commission_derived(row)
        for c in COMPUTED:
            gv = row.get(c)
            cv = got[c]
            if c == "AB":
                if str(gv).strip() != str(cv).strip():
                    fails[c].append((r, gv, cv))
            else:
                gvn = par._num(gv); cvn = par._num(cv)
                # relative + absolute tolerance for money columns
                tol = 1e-6 + 1e-4 * max(abs(gvn), abs(cvn))
                if abs(gvn - cvn) > tol:
                    fails[c].append((r, gvn, cvn))

    print(f"Rows checked: {n}\n")
    print(f"{'col':>3} {'label':32} {'mismatches':>10}")
    for c in COMPUTED:
        lab = str(ws.cell(row=3, column=ci(c)).value)[:32]
        print(f"{c:>3} {lab:32} {len(fails[c]):>10}")
    print()
    total = sum(len(v) for v in fails.values())
    if total == 0:
        print("=== ENGINE MATCHES GOLDEN ON ALL COMPUTED COLUMNS ===")
    else:
        print(f"=== {total} cell mismatches; samples: ===")
        for c in COMPUTED:
            if fails[c]:
                print(f"  {c}: " + "; ".join(f"row{r}: gold={g} got={k}" for r, g, k in fails[c][:4]))


if __name__ == "__main__":
    main()
